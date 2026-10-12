"""Real PostgreSQL tests for modeling governance, stable sync, and tenant isolation."""

import os
from io import BytesIO
from uuid import uuid4

import httpx
import pytest
from fastapi import FastAPI, Request
from openpyxl import Workbook

from mate_kernel.ontology.identity import ClassRef
from mate_platform.tenancy.context import AuthMethod, RequestContext
from mate_tech_ont.modeling.api import router
from mate_tech_ont.modeling.store import ModelingStore
from mate_tech_ont.v2_kernel import api as native
from mate_tech_ont.v2_kernel.pg_repo import PgOntologyRepository

DSN = os.environ.get("MODELING_TEST_DSN", "")
pytestmark = pytest.mark.skipif(not DSN, reason="MODELING_TEST_DSN requires isolated PostgreSQL")


@pytest.mark.asyncio
async def test_cross_domain_slug_collision_is_row_issue_and_can_be_resolved_by_reuse(app):
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        source = await candidate(client, resolutions={"item": {"mode": "new", "propertyMap": {}}})
        prefix = "/api/v1/ont/v2/modeling/sources/" + source["id"]
        checked = (
            await client.post(
                prefix + "/validate",
                json={"revision": source["revision"], "preparePublication": True},
            )
        ).json()
        native_type = checked["plan"]["types"][0]
        body = {
            **native_type["payload"],
            "rid": native_type["rid"].replace(".obj.data.", ".obj.scm."),
        }
        created = await client.post("/api/v1/ont/v2/object-types", json=body)
        assert created.status_code in (200, 201), created.text
        blocked = (
            await client.post(
                prefix + "/validate",
                json={"revision": source["revision"], "preparePublication": True},
            )
        ).json()
        assert not blocked["validation"]["canPublish"] and blocked["plan"] is None
        assert any(
            i["item"] == "item" and body["rid"] in i["message"]
            for i in blocked["validation"]["issues"]
        )
        source = (
            await client.put(
                prefix,
                json={
                    "revision": source["revision"],
                    "draft": source["draft"],
                    "resolutions": {
                        "item": {
                            "mode": "reuse",
                            "targetRid": body["rid"],
                            "propertyMap": native_type["propertyMap"],
                        }
                    },
                },
            )
        ).json()
        preview = (
            await client.post(prefix + "/completion-plan", json={"revision": source["revision"]})
        ).json()
        assert preview["canApply"], preview
        confirmed = await client.post(
            prefix + "/completion-confirm",
            json={"revision": source["revision"], "planHash": preview["hash"]},
        )
        assert confirmed.status_code == 200, confirmed.text
        confirmed_source = confirmed.json()
        prepared = (
            await client.post(
                prefix + "/validate",
                json={"revision": confirmed_source["revision"], "preparePublication": True},
            )
        ).json()
        assert prepared["plan"]["types"][0]["mode"] == "reuse"
        assert prepared["plan"]["types"][0]["rid"] == body["rid"]
        assert confirmed.json()["publication"] is None


@pytest.mark.asyncio
async def test_completion_native_collision_does_not_block_candidate_review(app):
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        source = await candidate(client)
        prefix = "/api/v1/ont/v2/modeling/sources/" + source["id"]
        body = {"revision": source["revision"]}
        plan = (await client.post(prefix + "/completion-plan", json=body)).json()
        assert plan["canApply"]
        other = await candidate(client)
        await validate_publish(client, other)
        updated = (await client.post(prefix + "/completion-plan", json=body)).json()
        assert updated["canApply"] and updated["hash"] == plan["hash"]
        confirmed = await client.post(
            prefix + "/completion-confirm", json={**body, "planHash": plan["hash"]}
        )
        assert confirmed.status_code == 200, confirmed.text
        result = confirmed.json()
        assert result["draft"]["classes"][0]["accepted"]
        assert result["validation"]["canPublish"] and result["plan"] is None
        assert result["publication"] is None
        checked = (
            await client.post(prefix + "/validate", json={"revision": result["revision"]})
        ).json()
        assert checked["validation"]["canPublish"] and checked["plan"] is None
        blocked = (
            await client.post(
                prefix + "/validate",
                json={"revision": result["revision"], "preparePublication": True},
            )
        ).json()
        assert not blocked["validation"]["canPublish"] and blocked["plan"] is None


@pytest.mark.asyncio
async def test_completion_plan_is_inert_hash_gated_and_deferred_class_never_syncs(app):
    from .test_modeling_generation_quality import related_fixture

    value, _ = related_fixture()
    book = Workbook()
    sheet = book.active
    sheet.title = "费用"
    sheet.append(["单号", "明细号", "金额"])
    sheet.append(["A", "10", 100])
    sheet.append(["B", "10", 200])
    payment = book.create_sheet("付款")
    payment.append(["付款号", "单号", "明细号"])
    payment.append(["P1", "A", "10"])
    output = BytesIO()
    book.save(output)
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        source = (
            await client.post(
                "/api/v1/ont/v2/modeling/sources",
                files={"file": ("completion.xlsx", output.getvalue())},
            )
        ).json()
        prefix = "/api/v1/ont/v2/modeling/sources/" + source["id"]
        source = (
            await client.put(prefix, json={"revision": source["revision"], "draft": value})
        ).json()
        body = {"revision": source["revision"], "deferClasses": ["line"]}
        plan = (await client.post(prefix + "/completion-plan", json=body)).json()
        assert (
            plan["canApply"]
            and plan["counts"]["classes"] == 1
            and plan["counts"]["relationships"] == 0
        )
        unchanged = (await client.get(prefix)).json()
        assert unchanged["draft"] == source["draft"] and unchanged["revision"] == source["revision"]
        assert "draft" not in plan
        assert (
            await client.post(prefix + "/completion-confirm", json={**body, "planHash": "0" * 64})
        ).status_code == 409
        assert (
            await client.post(
                prefix + "/completion-confirm",
                json={**body, "planHash": plan["hash"], "deferClasses": []},
            )
        ).status_code == 409
        result = await client.post(
            prefix + "/completion-confirm", json={**body, "planHash": plan["hash"]}
        )
        assert result.status_code == 200, result.text
        confirmed = result.json()
        assert (
            confirmed["revision"] == source["revision"] + 1
            and confirmed["validation"]["canPublish"]
            and confirmed["plan"] is None
        )
        assert len(confirmed["completionAudit"]) == 1 and confirmed["completionAudit"][0]["actor"]
        assert not confirmed["publication"] and not confirmed["sync"]
        assert (
            await client.post(
                prefix + "/completion-confirm", json={**body, "planHash": plan["hash"]}
            )
        ).status_code == 409
        published = await validate_publish(client, confirmed)
        synced = await client.post(prefix + "/sync", json={"revision": published["revision"]})
        assert synced.status_code == 200, synced.text
        assert synced.json()["sync"]["links"] == 0
        assert (
            len(published["plan"]["types"]) == 1
            and published["plan"]["types"][0]["classId"] == "payment"
        )
        repo = app.state.kernel_repo
        with repo.tenant_scope(app.state.test_tenant):
            instances = repo.list_individuals(ClassRef(published["plan"]["types"][0]["rid"]))
            assert len(instances) == 1
        assert (
            await client.post(prefix + "/completion-plan", json={"revision": published["revision"]})
        ).status_code == 422


async def treatment_candidate(client):
    from .test_modeling_data_treatment import source_fixture

    value, _ = source_fixture()
    book = Workbook()
    sheet = book.active
    sheet.title = "费用"
    sheet.append(["单号", "明细号", "金额"])
    for row in [["A", "10", 100], ["A", "10", 999], ["B", "10", 200], ["C", None, 300]]:
        sheet.append(row)
    pay = book.create_sheet("付款")
    pay.append(["付款号", "单号", "明细号"])
    pay.append(["P1", "A", "10"])
    pay.append(["P2", "B", "10"])
    output = BytesIO()
    book.save(output)
    source = (
        await client.post(
            "/api/v1/ont/v2/modeling/sources", files={"file": ("费用.xlsx", output.getvalue())}
        )
    ).json()
    prefix = f"/api/v1/ont/v2/modeling/sources/{source['id']}"
    source = (
        await client.put(prefix, json={"revision": source["revision"], "draft": value["draft"]})
    ).json()
    return await validate_publish(client, source)


@pytest.mark.asyncio
async def test_partial_sync_is_explicit_hash_gated_and_excludes_bad_objects_and_edges(app):
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        source = await treatment_candidate(client)
        prefix = f"/api/v1/ont/v2/modeling/sources/{source['id']}"
        assert (
            await client.post(prefix + "/sync", json={"revision": source["revision"]})
        ).status_code == 422
        checked = (
            await client.post(prefix + "/sync-check", json={"revision": source["revision"]})
        ).json()
        check = checked["syncCheck"]
        partial = check["partial"]
        assert (
            not check["canSync"]
            and partial["canSync"]
            and partial["preview"] == {"nodes": 3, "edges": 1}
        )
        assert partial["quarantine"]["rows"] == 3
        assert (
            await client.post(
                prefix + "/sync", json={"revision": source["revision"], "allowPartial": True}
            )
        ).status_code == 409
        exported = await client.get(
            prefix + "/quarantine", params={"checkHash": partial["checkHash"]}
        )
        assert (
            exported.status_code == 200 and "费用,2" in exported.text and "费用,5" in exported.text
        )
        assert (
            await client.get(prefix + "/quarantine", params={"checkHash": "0" * 64})
        ).status_code == 409
        body = {
            "revision": source["revision"],
            "allowPartial": True,
            "checkHash": partial["checkHash"],
        }
        # Emulate a prior failed run that staged A before writing any native objects.
        # The eligible subset must not accidentally import this old staged conflict.
        from copy import deepcopy

        from mate_tech_ont.modeling import data_treatment, ingestion

        store = ModelingStore(app.state.kernel_repo, app.state.test_tenant)
        _, rows = store.get(source["id"], private=True)
        key = data_treatment.candidates(source, rows)[0]["key"]
        proposed = data_treatment.proposal(
            source,
            rows,
            {"repairs": [{"key": key, "values": {"amount": 999}, "reason": "模拟旧批次暂存"}]},
        )
        trial = deepcopy(source)
        trial["dataTreatment"] = data_treatment.confirm(source, rows, proposed, [key], "test")
        _, old_nodes, _, _ = data_treatment.prepare(trial, rows, True)
        ingestion.stage(store, source, old_nodes)
        with ingestion.ingestion_lock(store):
            blocked = await client.post(prefix + "/sync", json=body)
            assert blocked.status_code == 422 and "已有 Excel 数据同步" in blocked.text
        for _ in range(2):
            result = await client.post(prefix + "/sync", json=body)
            assert result.status_code == 200, result.text
            assert result.json()["sync"]["status"] == "succeeded", result.text
            assert result.json()["sync"]["links"] == 1
            assert result.json()["sync"]["quarantine"]["rows"] == 3
        repo, tenant = app.state.kernel_repo, app.state.test_tenant
        with repo.tenant_scope(tenant):
            counts = [
                len(repo.list_individuals(ClassRef(t["rid"]))) for t in source["plan"]["types"]
            ]
            assert sorted(counts) == [1, 2]
        original, rows = ModelingStore(repo, tenant).get(source["id"], private=True)
        assert rows[0][0]["values"]["金额"] == 100 and original["draft"] == source["draft"]


@pytest.mark.asyncio
async def test_ai_proposal_confirmation_preserves_native_plan_and_invalidates_subset_authorization(
    app, monkeypatch
):
    import asyncio

    from mate_tech_ont.modeling import assistant_client

    async def propose(operation, payload, *args):
        assert operation == "data-repair"
        candidate = payload["validation"]["candidates"][0]
        return {
            "repairs": [
                {
                    "key": candidate["key"],
                    "values": {"amount": 999},
                    "reason": "来源修订值为999，请业务核对后确认",
                }
            ]
        }

    monkeypatch.setattr(assistant_client, "_post", propose)
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        source = await treatment_candidate(client)
        prefix = f"/api/v1/ont/v2/modeling/sources/{source['id']}"
        checked = (
            await client.post(prefix + "/sync-check", json={"revision": source["revision"]})
        ).json()
        old_hash = checked["syncCheck"]["partial"]["checkHash"]
        result = await client.post(prefix + "/data-repair", json={"revision": source["revision"]})
        assert result.status_code == 202
        assert result.json()["dataRepair"]["status"] == "running"
        await asyncio.gather(*app.state.modeling_tasks)
        proposed = (await client.get(prefix)).json()
        assert proposed["dataRepair"]["status"] == "succeeded", proposed["dataRepair"]
        assert not proposed.get("dataTreatment") and proposed["draft"] == source["draft"]
        proposal = proposed["dataRepair"]["proposal"]
        assert (
            proposal["after"]["counts"]["nodes"] == 4
            and proposal["after"]["quarantine"]["rows"] == 1
        )
        confirmed = await client.post(
            prefix + "/data-repair-confirm",
            json={
                "revision": source["revision"],
                "proposalId": proposal["id"],
                "keys": [proposal["repairs"][0]["key"]],
            },
        )
        assert confirmed.status_code == 200, confirmed.text
        confirmed = confirmed.json()
        assert (
            confirmed["plan"] == source["plan"]
            and confirmed["publication"] == source["publication"]
        )
        assert confirmed["syncCheck"] is None and confirmed["dataTreatment"]["audit"][0]["actor"]
        assert (
            await client.post(
                prefix + "/sync",
                json={
                    "revision": confirmed["revision"],
                    "allowPartial": True,
                    "checkHash": old_hash,
                },
            )
        ).status_code == 409
        checked = (
            await client.post(prefix + "/sync-check", json={"revision": confirmed["revision"]})
        ).json()
        result = await client.post(
            prefix + "/sync",
            json={
                "revision": confirmed["revision"],
                "allowPartial": True,
                "checkHash": checked["syncCheck"]["partial"]["checkHash"],
            },
        )
        assert result.json()["sync"]["status"] == "succeeded", result.text
        assert (
            result.json()["sync"]["links"] == 2 and result.json()["sync"]["eligible"]["nodes"] == 4
        )
        # Proposal cannot be reused after a confirmation changes effective values.
        assert (
            await client.post(
                prefix + "/data-repair-confirm",
                json={
                    "revision": confirmed["revision"],
                    "proposalId": proposal["id"],
                    "keys": [proposal["repairs"][0]["key"]],
                },
            )
        ).status_code == 422


@pytest.mark.asyncio
async def test_invalid_ai_values_fail_closed_and_restart_releases_abandoned_repair(
    app, monkeypatch
):
    import asyncio

    from mate_tech_ont.modeling import assistant_client

    async def propose(operation, payload, *args):
        return {
            "repairs": [
                {
                    "key": payload["validation"]["candidates"][0]["key"],
                    "values": {"amount": 123456},
                    "reason": "未出现在来源的值",
                }
            ]
        }

    monkeypatch.setattr(assistant_client, "_post", propose)
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        source = await treatment_candidate(client)
        prefix = f"/api/v1/ont/v2/modeling/sources/{source['id']}"
        await client.post(prefix + "/data-repair", json={"revision": source["revision"]})
        await asyncio.gather(*app.state.modeling_tasks)
        failed = (await client.get(prefix)).json()
        assert (
            failed["dataRepair"]["status"] == "failed"
            and not failed.get("dataTreatment")
            and failed["draft"] == source["draft"]
        )
        store = ModelingStore(app.state.kernel_repo, app.state.test_tenant)
        store.update(
            source["id"], source["revision"], {"dataRepair": {"status": "running"}}, increment=False
        )
        app.state.modeling_initialized.clear()
        restarted = (await client.get(prefix)).json()
        assert restarted["dataRepair"]["status"] == "failed"


@pytest.mark.asyncio
async def test_generation_baseline_is_stable_across_reimports_and_ignores_human_edits(
    app, monkeypatch
):
    import asyncio

    from mate_tech_ont.modeling import assistant_client

    calls = []

    async def propose(*args):
        calls.append(args)
        value = draft()
        for cls in value["classes"]:
            cls.update(
                accepted=False, identityStatus="proposed", identityReason="来源业务编码标识对象"
            )
        return value

    monkeypatch.setattr(assistant_client, "generate", propose)
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:

        async def generate_source(source):
            prefix = f"/api/v1/ont/v2/modeling/sources/{source['id']}"
            result = await client.post(
                prefix + "/generate",
                json={"revision": source["revision"], "context": "同一业务定义"},
            )
            assert result.status_code == 202, result.text
            await asyncio.gather(*app.state.modeling_tasks)
            value = (await client.get(prefix)).json()
            assert value["job"]["status"] == "succeeded", value["job"]
            return value

        first = await generate_source(await candidate(client))
        from copy import deepcopy

        baseline = deepcopy(first["draft"])
        assert first["generation"]["origin"] == "generated"
        first["draft"]["classes"][0]["name"] = "人工改名"
        saved = (
            await client.put(
                f"/api/v1/ont/v2/modeling/sources/{first['id']}",
                json={
                    "revision": first["revision"],
                    "draft": first["draft"],
                    "context": "同一业务定义",
                },
            )
        ).json()
        assert saved["generation"]["stale"] is True
        restored = await generate_source(saved)
        imported = await generate_source(await candidate(client))
        assert restored["draft"] == imported["draft"] == baseline
        assert imported["generation"]["origin"] == "reused" and len(calls) == 1
        assert calls[0][0]["platformCapabilities"]["identityFields"] == 4
        cache = ModelingStore(app.state.kernel_repo, app.state.test_tenant)
        fingerprint = first["generation"]["fingerprint"]
        other = ModelingStore(app.state.kernel_repo, "other-" + str(uuid4()))
        assert other.generation_baseline(fingerprint) is None
        other.generation_baseline(fingerprint, {"draft": {"name": "其他租户"}})
        assert cache.generation_baseline(fingerprint)["draft"] == baseline
        assert (
            cache.generation_baseline(fingerprint, {"draft": {"name": "不能覆盖"}})["draft"]
            == baseline
        )


@pytest.mark.asyncio
async def test_full_data_identity_repair_publishes_native_composite_identity_without_auto_approval(
    app, monkeypatch
):
    import asyncio

    from mate_tech_ont.modeling import assistant_client

    from .test_modeling_generation_quality import fixture

    value, _ = fixture()
    book = Workbook()
    sheet = book.active
    sheet.title = "费用"
    sheet.append(["单号", "明细号", "金额"])
    sheet.append(["A", "10", 100])
    sheet.append(["B", "10", 200])
    out = BytesIO()
    book.save(out)

    async def propose(*args):
        return value

    calls = []

    async def refine(operation, payload, *args):
        assert operation == "refine"
        calls.append(payload)
        assert payload["validation"]["identities"][0]["conflictingIdentities"] == 1
        return {
            "repairs": [
                {
                    "kind": "identity",
                    "id": "line",
                    "identity": ["line_no", "document_id"],
                    "reason": "明细编号在所属单据内唯一，需结合单据编号定位",
                }
            ]
        }

    monkeypatch.setattr(assistant_client, "generate", propose)
    monkeypatch.setattr(assistant_client, "_post", refine)
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        source = (
            await client.post(
                "/api/v1/ont/v2/modeling/sources", files={"file": ("expenses.xlsx", out.getvalue())}
            )
        ).json()
        prefix = f"/api/v1/ont/v2/modeling/sources/{source['id']}"
        assert (
            await client.post(prefix + "/generate", json={"revision": source["revision"]})
        ).status_code == 202
        await asyncio.gather(*app.state.modeling_tasks)
        source = (await client.get(prefix)).json()
        assert source["job"]["status"] == "succeeded", source["job"]
        cls = source["draft"]["classes"][0]
        assert (
            cls["identity"] == ["line_no", "document_id"]
            and cls["identityStatus"] == "proposed"
            and not cls["accepted"]
        )
        assert (
            source["generation"]["dataErrors"] == 0 and source["generation"]["structureErrors"] == 0
        )
        assert len(source["generation"]["changes"]) == 1 and len(calls) == 1
        assert source["publication"] is None and source["sync"] is None
        cls.update(accepted=True, identityStatus="confirmed")
        source["draft"]["mappings"][0]["accepted"] = True
        source = (
            await client.put(
                prefix, json={"revision": source["revision"], "draft": source["draft"]}
            )
        ).json()
        published = await validate_publish(client, source)
        rid = published["plan"]["types"][0]["rid"]
        with app.state.kernel_repo.tenant_scope(app.state.test_tenant):
            native_type = app.state.kernel_repo.get_object_type(ClassRef(rid))
            assert len(native_type.primary_key) == 2
        synchronized = await client.post(prefix + "/sync", json={"revision": published["revision"]})
        assert synchronized.status_code == 200, synchronized.text
        assert synchronized.json()["sync"]["status"] == "succeeded"
        assert synchronized.json()["sync"]["types"][0]["total_synced"] == 2
        assert (
            await client.post(prefix + "/generate", json={"revision": published["revision"]})
        ).status_code == 409


@pytest.mark.asyncio
async def test_pending_relation_saves_reloads_and_defers_without_publishing(app):
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        source = await candidate(client)
        prefix = f"/api/v1/ont/v2/modeling/sources/{source['id']}"
        value = source["draft"]
        ident = value["classes"][0]["id"]
        value["relationships"] = [
            {
                "id": "pending",
                "name": "待完善关联",
                "source": ident,
                "target": ident,
                "joins": [],
                "accepted": False,
                "targetSelection": "entity",
                "cardinalityStatus": "proposed",
            }
        ]
        saved = await client.put(prefix, json={"revision": source["revision"], "draft": value})
        assert saved.status_code == 200, saved.text
        restored = (await client.get(prefix)).json()
        assert restored["draft"]["relationships"][0]["joins"] == []
        checked = (
            await client.post(
                prefix + "/validate",
                json={"revision": restored["revision"], "preparePublication": True},
            )
        ).json()
        assert not checked["validation"]["canPublish"]
        assert any("关联字段映射" in issue["message"] for issue in checked["validation"]["issues"])
        value = restored["draft"]
        value["relationships"][0]["included"] = False
        deferred = (
            await client.put(prefix, json={"revision": restored["revision"], "draft": value})
        ).json()
        checked = (
            await client.post(
                prefix + "/validate",
                json={"revision": deferred["revision"], "preparePublication": True},
            )
        ).json()
        assert checked["validation"]["canPublish"]
        assert checked["publication"] is None and checked["sync"] is None


@pytest.mark.asyncio
async def test_sync_check_preserves_publication_and_async_sync_is_readable_and_idempotent(
    app, monkeypatch
):
    import asyncio
    import threading

    from mate_tech_ont.modeling import ingestion

    entered = threading.Event()
    release = threading.Event()
    original = ingestion.sync

    async def controlled(*args):
        entered.set()
        await asyncio.to_thread(release.wait, 10)
        return await original(*args)

    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        source = await validate_publish(client, await candidate(client))
        prefix = f"/api/v1/ont/v2/modeling/sources/{source['id']}"
        checked = await client.post(prefix + "/sync-check", json={"revision": source["revision"]})
        assert checked.status_code == 200, checked.text
        value = checked.json()
        assert value["plan"] == source["plan"] and value["publication"] == source["publication"]
        assert value["syncCheck"]["canSync"]
        assert value["syncCheck"]["preview"] == {"nodes": 2, "edges": 0}
        rid = source["plan"]["types"][0]["rid"]
        with app.state.kernel_repo.tenant_scope(app.state.test_tenant):
            assert not app.state.kernel_repo.list_individuals(
                ClassRef(rid), tenant_id=app.state.test_tenant
            )
        monkeypatch.setattr(ingestion, "sync", controlled)
        started = await client.post(
            prefix + "/sync", json={"revision": source["revision"], "background": True}
        )
        assert started.status_code == 202, started.text
        assert started.json()["sync"]["status"] == "running"
        try:
            assert await asyncio.to_thread(entered.wait, 5)
            assert (await client.get(prefix)).json()["sync"]["status"] == "running"
            duplicate = await client.post(
                prefix + "/sync", json={"revision": source["revision"], "background": True}
            )
            assert duplicate.status_code == 409
        finally:
            release.set()
        await asyncio.gather(*app.state.modeling_tasks)
        finished = (await client.get(prefix)).json()
        assert finished["sync"]["status"] == "succeeded", finished["sync"]
        assert finished["sync"]["types"][0]["total_synced"] == 2
        assert finished["sync"]["finishedAt"]
        again = await client.post(prefix + "/sync", json={"revision": source["revision"]})
        assert again.json()["sync"]["status"] == "succeeded"
        with app.state.kernel_repo.tenant_scope(app.state.test_tenant):
            assert (
                len(
                    app.state.kernel_repo.list_individuals(
                        ClassRef(rid), tenant_id=app.state.test_tenant
                    )
                )
                == 2
            )


@pytest.mark.asyncio
async def test_sync_check_blocks_missing_identity_and_native_drift_without_writes(app):
    book = Workbook()
    sheet = book.active
    sheet.title = "Items"
    sheet.append(["Code", "Name"])
    sheet.append([None, "First"])
    out = BytesIO()
    book.save(out)
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        source = await validate_publish(client, await candidate(client, out.getvalue()))
        prefix = f"/api/v1/ont/v2/modeling/sources/{source['id']}"
        checked = (
            await client.post(prefix + "/sync-check", json={"revision": source["revision"]})
        ).json()
        assert not checked["syncCheck"]["canSync"]
        assert checked["syncCheck"]["preview"] is None
        assert checked["syncCheck"]["report"]["errors"] == 1
        assert (
            await client.post(
                prefix + "/sync", json={"revision": source["revision"], "background": True}
            )
        ).status_code == 422
        rid = source["plan"]["types"][0]["rid"]
        with app.state.kernel_repo.tenant_scope(app.state.test_tenant):
            from dataclasses import replace

            original = app.state.kernel_repo.get_object_type(ClassRef(rid))
            app.state.kernel_repo.upsert_object_type(
                replace(original, description="Changed business definition")
            )
        checked = (
            await client.post(prefix + "/sync-check", json={"revision": source["revision"]})
        ).json()
        assert checked["syncCheck"]["baselineError"]


@pytest.mark.asyncio
async def test_sync_partial_relationship_failure_can_retry_without_duplicates(app, monkeypatch):
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        source = await candidate(client)
        from copy import deepcopy

        value = draft()
        target = deepcopy(value["classes"][0])
        target.update(id="target", name="Target")
        value["classes"].append(target)
        mapping = deepcopy(value["mappings"][0])
        mapping.update(id="target_rows", classId="target")
        value["mappings"].append(mapping)
        value["relationships"] = [
            {
                "id": "self",
                "name": "对应",
                "source": "item",
                "target": "target",
                "joins": [{"sourceProperty": "code", "targetProperty": "code"}],
                "accepted": True,
                "cardinality": "1:1",
                "cardinalityStatus": "confirmed",
                "cardinalityReason": "同一业务编码定位自身",
                "targetSelection": "entity",
            }
        ]
        prefix = f"/api/v1/ont/v2/modeling/sources/{source['id']}"
        source = (
            await client.put(prefix, json={"revision": source["revision"], "draft": value})
        ).json()
        source = await validate_publish(client, source)
        original = native.create_link_instance
        calls = 0

        async def fail_second(*args, **kwargs):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise ValueError("test relationship failure")
            return await original(*args, **kwargs)

        monkeypatch.setattr(native, "create_link_instance", fail_second)
        response = await client.post(prefix + "/sync", json={"revision": source["revision"]})
        assert response.json()["sync"]["status"] == "failed"
        assert response.json()["sync"]["links"] == 1, response.json()["sync"].get("error")
        assert response.json()["sync"]["types"][0]["total_synced"] == 2
        monkeypatch.setattr(native, "create_link_instance", original)
        response = await client.post(prefix + "/sync", json={"revision": source["revision"]})
        assert response.json()["sync"]["status"] == "succeeded"
        assert response.json()["sync"]["links"] == 2
        with app.state.kernel_repo.tenant_scope(app.state.test_tenant):
            conn, _ = app.state.kernel_repo._connect()
            with conn.cursor() as cur:
                cur.execute(
                    "SELECT count(*) FROM ont_link_instance WHERE tenant_id=%s",
                    (app.state.test_tenant,),
                )
                assert cur.fetchone()[0] == 2
            conn.close()


@pytest.mark.asyncio
async def test_sync_check_does_not_overwrite_changed_native_source_mapping(app):
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        source = await validate_publish(client, await candidate(client))
        prefix = f"/api/v1/ont/v2/modeling/sources/{source['id']}"
        response = await client.post(prefix + "/sync", json={"revision": source["revision"]})
        assert response.json()["sync"]["status"] == "succeeded"
        with app.state.kernel_repo.tenant_scope(app.state.test_tenant):
            repo = app.state.kernel_repo
            declaration = repo.list_backing_datasources(source["plan"]["types"][0]["rid"])[0]
            conn, _ = repo._connect()
            with conn.cursor() as cur:
                cur.execute(
                    "UPDATE ont_backing_datasource SET pk_column='changed_key' WHERE rid=%s",
                    (declaration["rid"],),
                )
            conn.commit()
            conn.close()
        checked = await client.post(prefix + "/sync-check", json={"revision": source["revision"]})
        assert not checked.json()["syncCheck"]["canSync"]
        assert "映射" in checked.json()["syncCheck"]["baselineError"]
        refused = await client.post(
            prefix + "/sync", json={"revision": source["revision"], "background": True}
        )
        assert refused.status_code == 422
        with repo.tenant_scope(app.state.test_tenant):
            assert (
                repo.list_backing_datasources(source["plan"]["types"][0]["rid"])[0]["pk_column"]
                == "changed_key"
            )


def workbook(name="First"):
    book = Workbook()
    sheet = book.active
    sheet.title = "Items"
    sheet.append(["Code", "Name"])
    sheet.append(["A", name])
    sheet.append(["B", "Second"])
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def draft():
    return {
        "schemaVersion": 2,
        "name": "Items",
        "classes": [
            {
                "id": "item",
                "name": "Item",
                "properties": [
                    {"id": "code", "name": "Code", "dataType": "text"},
                    {"id": "name", "name": "Name", "dataType": "text"},
                ],
                "identity": ["code"],
                "labelProperty": "name",
                "accepted": True,
            }
        ],
        "mappings": [
            {
                "id": "items",
                "classId": "item",
                "sheet": "Items",
                "bindings": {"code": "Code", "name": "Name"},
                "accepted": True,
            }
        ],
        "relationships": [],
    }


@pytest.fixture
def app(monkeypatch):
    monkeypatch.setenv("ONT_MODELING_SOURCE_DSN", DSN)
    application = FastAPI()
    application.state.kernel_repo = PgOntologyRepository(DSN)
    application.state.test_tenant = "model_" + uuid4().hex[:12]

    @application.middleware("http")
    async def ctx(request: Request, call_next):
        request.state.ctx = RequestContext(
            request_id="test",
            trace_id="test",
            tenant_id=request.headers.get("x-test-tenant", application.state.test_tenant),
            user_id="admin",
            roles=frozenset({request.headers.get("x-test-role", "PLATFORM_ADMIN")}),
            permissions=frozenset(),
            auth_method=AuthMethod.USER,
        )
        return await call_next(request)

    application.include_router(router)
    application.include_router(native.router)
    return application


async def candidate(client, content=None, resolutions=None):
    r = await client.post(
        "/api/v1/ont/v2/modeling/sources", files={"file": ("items.xlsx", content or workbook())}
    )
    assert r.status_code == 201, r.text
    source = r.json()
    r = await client.put(
        f"/api/v1/ont/v2/modeling/sources/{source['id']}",
        json={"revision": source["revision"], "draft": draft(), "resolutions": resolutions or {}},
    )
    assert r.status_code == 200, r.text
    return r.json()


async def validate_publish(client, source):
    prefix = f"/api/v1/ont/v2/modeling/sources/{source['id']}"
    r = await client.post(
        prefix + "/validate", json={"revision": source["revision"], "preparePublication": True}
    )
    assert r.status_code == 200, r.text
    source = r.json()
    assert source["validation"]["canPublish"], source
    r = await client.post(
        prefix + "/publish",
        json={
            "revision": source["revision"],
            "configHash": source["plan"]["configHash"],
            "acknowledgeWarnings": True,
        },
    )
    assert r.status_code == 200, r.text
    assert r.json()["publication"]["status"] == "succeeded", r.text
    return r.json()


@pytest.mark.asyncio
async def test_deferred_scope_roundtrip_native_publish_sync_and_scope_lock(app):
    from copy import deepcopy

    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        source = await candidate(client)
        value = draft()
        later = deepcopy(value["classes"][0])
        later.update(id="later", included=False, accepted=False, identityStatus="proposed")
        value["classes"].append(later)
        value["relationships"] = [
            {
                "id": "later_link",
                "name": "Later",
                "source": "item",
                "target": "later",
                "included": False,
                "joins": [{"sourceProperty": "code", "targetProperty": "code"}],
                "cardinalityStatus": "proposed",
            }
        ]
        prefix = f"/api/v1/ont/v2/modeling/sources/{source['id']}"
        saved = await client.put(prefix, json={"revision": source["revision"], "draft": value})
        assert saved.status_code == 200, saved.text
        reopened = (await client.get(prefix)).json()
        assert len(reopened["draft"]["classes"]) == 2
        assert reopened["draft"]["classes"][1]["included"] is False
        published = await validate_publish(client, reopened)
        assert [t["classId"] for t in published["plan"]["types"]] == ["item"]
        assert published["plan"]["relationships"] == []
        synced = await client.post(prefix + "/sync", json={"revision": published["revision"]})
        assert synced.status_code == 200, synced.text
        assert synced.json()["sync"]["status"] == "succeeded"
        value["classes"][1]["included"] = True
        rejected = await client.put(
            prefix, json={"revision": published["revision"], "draft": value}
        )
        assert rejected.status_code == 409
        assert (await client.get(prefix)).json()["draft"]["classes"][1]["included"] is False


@pytest.mark.asyncio
async def test_modeling_cardinality_native_acceptance_and_read_only_drift_detection(app):
    from copy import deepcopy

    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        source = await candidate(client)
        value = draft()
        target = deepcopy(value["classes"][0])
        target.update(id="target", name="Target")
        value["classes"].append(target)
        mapping = deepcopy(value["mappings"][0])
        mapping.update(id="target_rows", classId="target")
        value["mappings"].append(mapping)
        value["relationships"] = [
            {
                "id": "belongs",
                "name": "belongs",
                "source": "item",
                "target": "target",
                "joins": [{"sourceProperty": "code", "targetProperty": "code"}],
                "accepted": True,
                "cardinality": "N:1",
                "cardinalityStatus": "confirmed",
                "cardinalityReason": "每条对象引用一个目标",
                "targetSelection": "entity",
            }
        ]
        prefix = f"/api/v1/ont/v2/modeling/sources/{source['id']}"
        saved = await client.put(prefix, json={"revision": source["revision"], "draft": value})
        source = await validate_publish(client, saved.json())
        relation = source["plan"]["relationships"][0]
        assert relation["payload"]["cardinality"] == "N:1"
        report = await client.get(prefix + "/acceptance")
        assert report.status_code == 200, report.text
        assert report.json()["status"] == "passed"
        assert source["publication"]["acceptance"]["status"] == "passed"
        changed = {**relation["payload"], "cardinality": "N:N"}
        response = await client.post("/api/v1/ont/v2/link-types", json=changed)
        assert response.status_code == 200, response.text
        report = (await client.get(prefix + "/acceptance")).json()
        assert report["status"] == "drifted"
        assert report["relationships"][0]["status"] == "changed"
        reread = (await client.get(prefix)).json()
        assert reread["revision"] == source["revision"]
        assert reread["publication"]["status"] == "succeeded"


@pytest.mark.asyncio
async def test_modeling_comparison_preserves_native_definition_and_candidates(app):
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        original = await validate_publish(client, await candidate(client))
        other = await candidate(client)
        rid = original["plan"]["types"][0]["rid"]
        prefix = f"/api/v1/ont/v2/modeling/sources/{other['id']}"
        response = await client.get(
            prefix + "/comparison", params={"classId": "item", "targetRid": rid}
        )
        assert response.status_code == 200, response.text
        result = response.json()
        assert result["classId"] == "item" and result["targetRid"] == rid
        assert result["requiresDecision"] and len(result["properties"]) == 2
        assert (await client.get(prefix)).json()["revision"] == other["revision"]


@pytest.mark.asyncio
async def test_modeling_roundtrip_stable_sync_and_cross_file_reuse(app):
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        source = await validate_publish(client, await candidate(client))
        prefix = f"/api/v1/ont/v2/modeling/sources/{source['id']}"
        for _ in range(2):
            r = await client.post(prefix + "/sync", json={"revision": source["revision"]})
            assert r.status_code == 200, r.text
            assert r.json()["sync"]["status"] == "succeeded", r.text
        native_type = source["plan"]["types"][0]
        repo = app.state.kernel_repo
        with repo.tenant_scope(app.state.test_tenant):
            first = repo.list_individuals(ClassRef(native_type["rid"]))
        assert len(first) == 2
        another = await candidate(
            client,
            workbook("Changed"),
            {
                "item": {
                    "mode": "reuse",
                    "targetRid": native_type["rid"],
                    "propertyMap": native_type["propertyMap"],
                }
            },
        )
        another = await validate_publish(client, another)
        r = await client.post(
            f"/api/v1/ont/v2/modeling/sources/{another['id']}/sync",
            json={"revision": another["revision"]},
        )
        assert r.json()["sync"]["status"] == "succeeded", r.text
        with repo.tenant_scope(app.state.test_tenant):
            second = repo.list_individuals(ClassRef(native_type["rid"]))
        assert {i.rid for i in first} == {i.rid for i in second}
        assert any(i.get(ClassRef(native_type["propertyMap"]["name"])) == "Changed" for i in second)


@pytest.mark.asyncio
async def test_modeling_revision_auth_and_tenant_isolation(app):
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        source = await candidate(client)
        prefix = f"/api/v1/ont/v2/modeling/sources/{source['id']}"
        assert (await client.get(prefix, headers={"x-test-tenant": "other"})).status_code == 404
        assert (
            await client.post(
                prefix + "/validate", json={"revision": 1, "preparePublication": True}
            )
        ).status_code == 409
        assert (
            await client.put(
                prefix,
                json={"revision": source["revision"], "draft": draft()},
                headers={"x-test-role": "viewer"},
            )
        ).status_code == 403
        body = (await client.get(prefix)).json()
        assert "rows" not in body and "file_data" not in body
        store = ModelingStore(app.state.kernel_repo, "other")
        assert not any(s["id"] == source["id"] for s in store.list())


@pytest.mark.asyncio
async def test_modeling_collision_requires_explicit_reuse_and_stale_baseline_rejected(app):
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        source = await validate_publish(client, await candidate(client))
        collision = await candidate(client)
        r = await client.post(
            f"/api/v1/ont/v2/modeling/sources/{collision['id']}/validate",
            json={"revision": collision["revision"], "preparePublication": True},
        )
        assert r.status_code == 200 and not r.json()["validation"]["canPublish"]
        assert r.json()["plan"] is None and "复用" in r.text
        native_type = source["plan"]["types"][0]
        reused = await candidate(
            client,
            resolutions={
                "item": {
                    "mode": "reuse",
                    "targetRid": native_type["rid"],
                    "propertyMap": native_type["propertyMap"],
                }
            },
        )
        prefix = f"/api/v1/ont/v2/modeling/sources/{reused['id']}"
        reused = (
            await client.post(
                prefix + "/validate",
                json={"revision": reused["revision"], "preparePublication": True},
            )
        ).json()
        repo = app.state.kernel_repo
        from dataclasses import replace

        with repo.tenant_scope(app.state.test_tenant):
            ot = repo.get_object_type(ClassRef(native_type["rid"]))
            repo.upsert_object_type(replace(ot, display_name="Concurrent edit"))
        r = await client.post(
            prefix + "/publish",
            json={
                "revision": reused["revision"],
                "configHash": reused["plan"]["configHash"],
                "acknowledgeWarnings": True,
            },
        )
        assert r.status_code == 422 and "修改" in r.text


@pytest.mark.asyncio
async def test_modeling_relations_extend_and_manual_overlay(app):
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        source = await candidate(client)
        value = draft()
        value["classes"].append(
            {
                "id": "label",
                "name": "Label",
                "properties": [{"id": "name", "name": "Name", "dataType": "text"}],
                "identity": ["name"],
                "labelProperty": "name",
                "accepted": True,
            }
        )
        value["mappings"].append(
            {
                "id": "labels",
                "classId": "label",
                "sheet": "Items",
                "bindings": {"name": "Name"},
                "accepted": True,
            }
        )
        value["relationships"] = [
            {
                "id": "has_label",
                "name": "Has label",
                "source": "item",
                "target": "label",
                "joins": [{"sourceProperty": "name", "targetProperty": "name"}],
                "accepted": True,
            }
        ]
        prefix = f"/api/v1/ont/v2/modeling/sources/{source['id']}"
        source = (
            await client.put(prefix, json={"revision": source["revision"], "draft": value})
        ).json()
        source = await validate_publish(client, source)
        r = await client.post(prefix + "/sync", json={"revision": source["revision"]})
        assert r.json()["sync"]["status"] == "succeeded", r.text
        assert r.json()["sync"]["links"] == 2
        r = await client.post(prefix + "/sync", json={"revision": source["revision"]})
        assert r.json()["sync"]["status"] == "succeeded", r.text
        native_type = source["plan"]["types"][0]
        repo, tenant = app.state.kernel_repo, app.state.test_tenant
        with repo.tenant_scope(tenant):
            inds = repo.list_individuals(ClassRef(native_type["rid"]))
            original = inds[0]
            name_rid = native_type["propertyMap"]["name"]
            from dataclasses import replace

            repo.create_individual(
                replace(
                    original,
                    props=tuple(
                        (p, "Manual" if p.rid == name_rid else v) for p, v in original.props
                    ),
                )
            )
        store = ModelingStore(repo, tenant)
        with store.connection() as cur:
            cur.execute(
                "INSERT INTO ont_edit_overlay(individual_rid,property_rid,tenant_id) VALUES (%s,%s,%s)",
                (original.rid, name_rid, tenant),
            )
        r = await client.post(prefix + "/sync", json={"revision": source["revision"]})
        assert r.json()["sync"]["status"] == "succeeded", r.text
        with repo.tenant_scope(tenant):
            assert repo.get_individual(original.rid).get(ClassRef(name_rid)) == "Manual"
        # Extend the existing type without changing its identity or deleting properties.
        extended = await candidate(
            client,
            resolutions={
                "item": {
                    "mode": "extend",
                    "targetRid": native_type["rid"],
                    "propertyMap": native_type["propertyMap"],
                }
            },
        )
        value = draft()
        value["classes"][0]["properties"].append(
            {"id": "alias", "name": "Alias", "dataType": "text"}
        )
        value["mappings"][0]["bindings"]["alias"] = "Name"
        extended = (
            await client.put(
                f"/api/v1/ont/v2/modeling/sources/{extended['id']}",
                json={
                    "revision": extended["revision"],
                    "draft": value,
                    "resolutions": extended["resolutions"],
                },
            )
        ).json()
        extended = await validate_publish(client, extended)
        with repo.tenant_scope(tenant):
            ot = repo.get_object_type(ClassRef(native_type["rid"]))
            assert len(ot.properties) == 3


@pytest.mark.asyncio
async def test_modeling_failed_generation_preserves_draft_and_restart_marks_interrupted(
    app, monkeypatch
):
    from mate_tech_ont.modeling import assistant_client

    async def fail(*args):
        raise ValueError("provider unavailable")

    monkeypatch.setattr(assistant_client, "generate", fail)
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        source = await candidate(client)
        prefix = f"/api/v1/ont/v2/modeling/sources/{source['id']}"
        r = await client.post(
            prefix + "/generate",
            json={"revision": source["revision"], "context": "Build item model", "questions": []},
        )
        assert r.status_code == 202
        import asyncio

        await asyncio.gather(*app.state.modeling_tasks)
        reopened = (await client.get(prefix)).json()
        assert reopened["job"]["status"] == "failed"
        assert reopened["draft"] == source["draft"]
        store = ModelingStore(app.state.kernel_repo, app.state.test_tenant)
        store.update(
            source["id"],
            reopened["revision"],
            {"job": {"id": "old", "kind": "generate", "status": "running"}},
            increment=False,
        )
        app.state.modeling_initialized.clear()
        reopened = (await client.get(prefix)).json()
        assert reopened["job"]["status"] == "interrupted"


@pytest.mark.asyncio
async def test_modeling_publication_preserves_existing_native_editor_wip(app):
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        original = await validate_publish(client, await candidate(client))
        native_type = original["plan"]["types"][0]
        repo = app.state.kernel_repo
        with repo.tenant_scope(app.state.test_tenant):
            native_wip = {**native_type["payload"], "description": "Unpublished human change"}
            repo.save_schema_wip(native_type["rid"], native_wip, "positive-flow-user")
        source = await candidate(
            client,
            resolutions={
                "item": {
                    "mode": "extend",
                    "targetRid": native_type["rid"],
                    "propertyMap": native_type["propertyMap"],
                }
            },
        )
        prefix = f"/api/v1/ont/v2/modeling/sources/{source['id']}"
        source = (
            await client.post(
                prefix + "/validate",
                json={"revision": source["revision"], "preparePublication": True},
            )
        ).json()
        r = await client.post(
            prefix + "/publish",
            json={
                "revision": source["revision"],
                "configHash": source["plan"]["configHash"],
                "acknowledgeWarnings": True,
            },
        )
        assert r.status_code == 200, r.text
        assert r.json()["publication"]["status"] == "failed"
        assert "草稿" in r.json()["publication"]["error"]
        with repo.tenant_scope(app.state.test_tenant):
            preserved = repo.get_schema_wip(native_type["rid"])
            assert preserved["payload"] == native_wip
            assert preserved["author"] == "positive-flow-user"


@pytest.mark.asyncio
async def test_modeling_reuse_updates_existing_positive_flow_instance(app):
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        source = await validate_publish(client, await candidate(client))
        item = source["plan"]["types"][0]
        tenant = app.state.test_tenant
        rid = f"ont.{tenant}.ind.item.A"
        r = await client.post(
            "/api/v1/ont/v2/individuals",
            json={
                "rid": rid,
                "class_rid": item["rid"],
                "primary_key": "A",
                "props": {
                    item["propertyMap"]["code"]: {"value": "A"},
                    item["propertyMap"]["name"]: {"value": "Positive flow"},
                },
            },
        )
        assert r.status_code == 200, r.text
        reused = await candidate(
            client,
            resolutions={
                "item": {
                    "mode": "reuse",
                    "targetRid": item["rid"],
                    "propertyMap": item["propertyMap"],
                }
            },
        )
        reused = await validate_publish(client, reused)
        r = await client.post(
            f"/api/v1/ont/v2/modeling/sources/{reused['id']}/sync",
            json={"revision": reused["revision"]},
        )
        assert r.json()["sync"]["status"] == "succeeded", r.text
        with app.state.kernel_repo.tenant_scope(tenant):
            instances = app.state.kernel_repo.list_individuals(ClassRef(item["rid"]))
        assert len(instances) == 2
        assert any(
            i.rid == rid and i.get(ClassRef(item["propertyMap"]["name"])) == "First"
            for i in instances
        )


@pytest.mark.asyncio
async def test_modeling_other_worker_does_not_interrupt_live_job(app, monkeypatch):
    import asyncio

    from mate_tech_ont.modeling import assistant_client

    release = asyncio.Event()
    reported = asyncio.Event()

    async def slow(*args):
        await assistant_client.generation_progress.get()(
            {"stage": "mapping", "batch": 2, "totalBatches": 9}
        )
        reported.set()
        await release.wait()
        return draft()

    monkeypatch.setattr(assistant_client, "generate", slow)
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        source = await candidate(client)
        prefix = f"/api/v1/ont/v2/modeling/sources/{source['id']}"
        r = await client.post(
            prefix + "/generate",
            json={
                "revision": source["revision"],
                "context": "Items",
                "questions": [{"id": "q1", "question": "Which items?", "expected": "List items"}],
            },
        )
        assert r.status_code == 202, r.text
        await asyncio.wait_for(reported.wait(), timeout=5)
        other_worker = ModelingStore(PgOntologyRepository(DSN), app.state.test_tenant)
        other_worker.interrupt_jobs()
        assert other_worker.get(source["id"])["job"]["status"] == "running"
        assert other_worker.get(source["id"])["job"]["progress"] == {
            "stage": "mapping",
            "batch": 2,
            "totalBatches": 9,
        }
        assert other_worker.get(source["id"])["revision"] == r.json()["revision"]
        assert other_worker.get(source["id"])["questions"][0]["id"] == "q1"
        release.set()
        await asyncio.gather(*app.state.modeling_tasks)
        assert other_worker.get(source["id"])["job"]["status"] == "succeeded"


@pytest.mark.asyncio
@pytest.mark.parametrize("race_point", ["save", "apply", "delete"])
async def test_modeling_native_wip_concurrent_editor_is_preserved(app, monkeypatch, race_point):
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        original = await validate_publish(client, await candidate(client))
        item = original["plan"]["types"][0]
        source = await candidate(
            client,
            resolutions={
                "item": {
                    "mode": "extend",
                    "targetRid": item["rid"],
                    "propertyMap": item["propertyMap"],
                }
            },
        )
        prefix = f"/api/v1/ont/v2/modeling/sources/{source['id']}"
        source = (
            await client.post(
                prefix + "/validate",
                json={"revision": source["revision"], "preparePublication": True},
            )
        ).json()
        human = {**item["payload"], "description": "Concurrent native edit"}

        def edit():
            with app.state.kernel_repo.tenant_scope(app.state.test_tenant):
                app.state.kernel_repo.save_schema_wip(item["rid"], human, "human")

        if race_point == "save":
            original_fn = native.save_schema_wip

            async def race(*args, **kwargs):
                edit()
                return await original_fn(*args, **kwargs)

            monkeypatch.setattr(native, "save_schema_wip", race)
        elif race_point == "apply":
            original_fn = native.apply_schema_wip

            async def race(*args, **kwargs):
                edit()
                return await original_fn(*args, **kwargs)

            monkeypatch.setattr(native, "apply_schema_wip", race)
        else:
            original_fn = native._upsert_object_type_gated

            async def race(*args, **kwargs):
                result = await original_fn(*args, **kwargs)
                edit()
                return result

            monkeypatch.setattr(native, "_upsert_object_type_gated", race)
        r = await client.post(
            prefix + "/publish",
            json={
                "revision": source["revision"],
                "configHash": source["plan"]["configHash"],
                "acknowledgeWarnings": True,
            },
        )
        assert r.status_code == 200, r.text
        assert r.json()["publication"]["status"] == (
            "succeeded" if race_point == "delete" else "failed"
        ), r.text
        with app.state.kernel_repo.tenant_scope(app.state.test_tenant):
            assert app.state.kernel_repo.get_schema_wip(item["rid"])["payload"] == human
            assert (
                app.state.kernel_repo.get_object_type(ClassRef(item["rid"])).description
                != human["description"]
            )


def equipment_hierarchy():
    code = {"id": "equipment_code", "name": "设备编码", "dataType": "text"}
    name = {"id": "name", "name": "设备名称", "dataType": "text"}
    capacity = {"id": "capacity", "name": "额定功率", "dataType": "number"}
    axes = {"id": "axes", "name": "轴数", "dataType": "number"}

    def cls(ident, title, props, parent=None):
        return {
            "id": ident,
            "name": title,
            "properties": props,
            "identity": ["equipment_code"],
            "labelProperty": "name",
            "accepted": True,
            "parentClassId": parent,
            "inheritanceReason": "业务分类：每一台数控机床都是机床，每一台机床都是设备"
            if parent
            else "",
        }

    return {
        "schemaVersion": 2,
        "name": "设备分类",
        "classes": [
            cls("cnc", "数控机床", [code, name, capacity, axes], "machine"),
            cls("equipment", "设备", [code, name]),
            cls("machine", "机床", [code, name, capacity], "equipment"),
        ],
        "mappings": [
            {
                "id": "cnc_rows",
                "classId": "cnc",
                "sheet": "Equipment",
                "bindings": {
                    "equipment_code": "Code",
                    "name": "Name",
                    "capacity": "Power",
                    "axes": "Axes",
                },
                "accepted": True,
            }
        ],
        "relationships": [],
    }


def equipment_content():
    book = Workbook()
    sheet = book.active
    sheet.title = "Equipment"
    sheet.append(["Code", "Name", "Power", "Axes"])
    sheet.append(["CNC-1", "机床一", 15, 3])
    sheet.append(["CNC-2", "机床二", 20, 5])
    content = BytesIO()
    book.save(content)
    return content.getvalue()


@pytest.mark.asyncio
async def test_modeling_three_level_native_hierarchy_and_leaf_only_instances(app):
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        r = await client.post(
            "/api/v1/ont/v2/modeling/sources",
            files={"file": ("equipment.xlsx", equipment_content())},
        )
        assert r.status_code == 201, r.text
        source = r.json()
        prefix = f"/api/v1/ont/v2/modeling/sources/{source['id']}"
        r = await client.put(
            prefix, json={"revision": source["revision"], "draft": equipment_hierarchy()}
        )
        assert r.status_code == 200, r.text
        source = await validate_publish(client, r.json())
        types = source["plan"]["types"]
        assert [t["classId"] for t in types] == ["equipment", "machine", "cnc"]
        assert types[1]["payload"]["parent_class"] == types[0]["rid"]
        assert types[2]["payload"]["parent_class"] == types[1]["rid"]
        assert len({t["propertyMap"]["equipment_code"] for t in types}) == 1
        r = await client.post(prefix + "/sync", json={"revision": source["revision"]})
        assert r.json()["sync"]["status"] == "succeeded", r.text
        assert [t["total_synced"] for t in r.json()["sync"]["types"]] == [0, 0, 2]
        r = await client.get("/api/v1/ont/v2/individuals", params={"class_rid": types[0]["rid"]})
        assert r.status_code == 200, r.text
        assert len(r.json()) == 2
        assert {i["class_rid"] for i in r.json()} == {types[2]["rid"]}
        for item in types:
            r = await client.get("/api/v1/ont/v2/object-types/" + item["rid"])
            assert r.status_code == 200, r.text
            assert r.json().get("parent_class", "") == item["payload"].get("parent_class", "")


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "invalid", ["missing", "self", "cycle", "identity", "property", "evidence"]
)
async def test_modeling_invalid_hierarchy_cannot_prepare_publication(app, invalid):
    value = equipment_hierarchy()
    cnc, equipment, machine = value["classes"]
    if invalid == "missing":
        cnc["parentClassId"] = "unknown"
    elif invalid == "self":
        cnc["parentClassId"] = "cnc"
    elif invalid == "cycle":
        equipment.update(parentClassId="cnc", inheritanceReason="Circular assertion")
    elif invalid == "identity":
        cnc["identity"] = ["name"]
    elif invalid == "property":
        cnc["properties"] = [p for p in cnc["properties"] if p["id"] != "capacity"]
        del value["mappings"][0]["bindings"]["capacity"]
    else:
        cnc["inheritanceReason"] = ""
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        r = await client.post(
            "/api/v1/ont/v2/modeling/sources",
            files={"file": ("equipment.xlsx", equipment_content())},
        )
        assert r.status_code == 201, r.text
        source = r.json()
        prefix = f"/api/v1/ont/v2/modeling/sources/{source['id']}"
        r = await client.put(prefix, json={"revision": source["revision"], "draft": value})
        assert r.status_code == 200, r.text
        source = r.json()
        r = await client.post(
            prefix + "/validate", json={"revision": source["revision"], "preparePublication": True}
        )
        assert r.status_code == 422 or (
            r.status_code == 200 and r.json()["validation"]["errors"] > 0
        ), r.text
        with app.state.kernel_repo.tenant_scope(app.state.test_tenant):
            assert (
                app.state.kernel_repo.list_object_types(100, 0, tenant_id=app.state.test_tenant)
                == []
            )


@pytest.mark.asyncio
async def test_structure_can_publish_with_missing_data_identity_but_sync_is_blocked(app):
    book = Workbook()
    sheet = book.active
    sheet.title = "Items"
    sheet.append(["Code", "Name"])
    sheet.append([None, "First"])
    sheet.append([None, "Second"])
    out = BytesIO()
    book.save(out)
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        source = await candidate(client, out.getvalue())
        source = await validate_publish(client, source)
        report = source["validation"]
        assert report["scope"] == "ontology"
        assert report["errors"] == 0 and report["canPublish"]
        assert report["dataValidation"]["errors"] == 2
        assert not report["dataValidation"]["canSync"]
        assert source["publication"]["status"] == "succeeded"
        prefix = f"/api/v1/ont/v2/modeling/sources/{source['id']}"
        result = await client.post(prefix + "/sync", json={"revision": source["revision"]})
        assert result.status_code == 422
        assert "数据" in result.text
        with app.state.kernel_repo.tenant_scope(app.state.test_tenant):
            types = app.state.kernel_repo.list_object_types(100, 0, tenant_id=app.state.test_tenant)
            assert len(types) == 1
            assert (
                app.state.kernel_repo.list_individuals(
                    ClassRef(source["plan"]["types"][0]["rid"]), tenant_id=app.state.test_tenant
                )
                == []
            )
