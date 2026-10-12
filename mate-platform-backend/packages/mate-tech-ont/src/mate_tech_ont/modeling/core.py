"""Bounded, data-driven modeling. Model proposals are inert configuration only."""

import hashlib
import json
import math
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Annotated, Literal
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from pydantic import (
    BaseModel,
    ConfigDict,
    Field,
    StrictBool,
    StrictFloat,
    StrictInt,
    StrictStr,
    ValidationError,
)

Slug = Annotated[str, Field(pattern=r"^[a-zA-Z][a-zA-Z0-9_-]{0,39}$")]
Name = Annotated[str, Field(min_length=1, max_length=120)]
EntityKey = Annotated[str, Field(min_length=1, max_length=200)]
Primitive = StrictStr | StrictInt | StrictFloat | StrictBool | None


class Strict(BaseModel):
    model_config = ConfigDict(extra="forbid", str_strip_whitespace=True, allow_inf_nan=False)


class Property(Strict):
    included: StrictBool = True
    deferReason: str = Field(default="", max_length=1000)
    id: Slug
    name: Name
    dataType: Literal["text", "number", "date"]


class Evidence(Strict):
    kind: Literal["source", "business", "inference"]
    detail: str = Field(min_length=1, max_length=1000)
    sheet: str = Field(default="", max_length=120)
    columns: list[Name] = Field(default_factory=list, max_length=12)


class UniqueKey(Strict):
    properties: list[Slug] = Field(min_length=1, max_length=4)
    reason: str = Field(default="", max_length=1000)
    confirmed: StrictBool = False


class Class(Strict):
    included: StrictBool = True
    id: Slug
    name: Name
    description: str = Field(default="", max_length=1000)
    properties: list[Property] = Field(min_length=1, max_length=40)
    identity: list[Slug] = Field(max_length=4)
    labelProperty: Slug
    accepted: StrictBool = False
    parentClassId: Slug | None = None
    inheritanceReason: str = Field(default="", max_length=1000)
    grain: str = Field(default="", max_length=500)
    identityStatus: Literal["legacy", "proposed", "confirmed", "unresolved", "source-record"] = (
        "legacy"
    )
    identityCandidate: list[Slug] = Field(default_factory=list, max_length=4)
    identityReason: str = Field(default="", max_length=1000)
    uniqueKeys: list[UniqueKey] = Field(default_factory=list, max_length=4)
    evidence: list[Evidence] = Field(default_factory=list, max_length=8)


class Mapping(Strict):
    id: Slug
    classId: Slug
    sheet: Name
    bindings: dict[Slug, Name]
    reason: str = Field(default="", max_length=1000)
    accepted: StrictBool = False


class Join(Strict):
    sourceProperty: Slug | Literal[""]
    targetProperty: Slug | Literal[""]


class Relationship(Strict):
    included: StrictBool = True
    id: Slug
    name: Name
    source: Slug
    target: Slug
    joins: list[Join] = Field(default_factory=list, max_length=4)
    reason: str = Field(default="", max_length=1000)
    accepted: StrictBool = False
    cardinality: Literal["1:1", "1:N", "N:1", "N:N"] = "N:N"
    cardinalityStatus: Literal["legacy", "proposed", "confirmed", "unresolved"] = "legacy"
    cardinalityReason: str = Field(default="", max_length=1000)
    targetSelection: Literal["legacy", "entity", "group"] = "legacy"
    evidence: list[Evidence] = Field(default_factory=list, max_length=8)


class ManualInstance(Strict):
    id: Slug
    classId: Slug
    values: dict[Slug, Primitive]
    accepted: StrictBool = False


class InstanceEdit(Strict):
    key: EntityKey
    values: dict[Slug, Primitive] = Field(default_factory=dict)
    excluded: StrictBool = False


class LinkEdit(Strict):
    id: Slug
    relationId: Slug
    source: EntityKey
    target: EntityKey
    excluded: StrictBool = False


class ModelingQuestion(Strict):
    id: Slug
    question: str = Field(min_length=1, max_length=500)
    expected: str = Field(default="", max_length=1000)


class SemanticDefinition(Strict):
    target: str = Field(
        pattern=r"^(?:(class|mapping|relationship):[a-zA-Z][a-zA-Z0-9_-]{0,39}|property:[a-zA-Z][a-zA-Z0-9_-]{0,39}:[a-zA-Z][a-zA-Z0-9_-]{0,39})$"
    )
    definition: str = Field(default="", max_length=1000)
    aliases: list[Name] = Field(default_factory=list, max_length=12)
    scope: str = Field(default="", max_length=500)
    grain: str = Field(default="", max_length=300)
    unit: str = Field(default="", max_length=120)
    constraints: str = Field(default="", max_length=1000)
    basis: str = Field(default="", max_length=1000)
    confirmed: StrictBool = False


class ModelingSemantics(Strict):
    context: str = Field(default="", max_length=2000)
    questions: list[ModelingQuestion] = Field(default_factory=list, max_length=20)
    definitions: list[SemanticDefinition] = Field(default_factory=list, max_length=150)


class Draft(Strict):
    schemaVersion: Literal[2] = 2
    name: Name
    classes: list[Class] = Field(min_length=1, max_length=30)
    mappings: list[Mapping] = Field(max_length=60)
    relationships: list[Relationship] = Field(max_length=60)
    instances: list[ManualInstance] = Field(default_factory=list, max_length=500)
    instanceEdits: list[InstanceEdit] = Field(default_factory=list, max_length=1000)
    links: list[LinkEdit] = Field(default_factory=list, max_length=1000)
    notes: list[Annotated[str, Field(max_length=1000)]] = Field(default_factory=list, max_length=30)
    semantics: ModelingSemantics = Field(default_factory=ModelingSemantics)


def cell(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def digest(value):
    return hashlib.sha256(
        json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()


def read_workbook(content):
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 80 * 1024 * 1024:
                raise ValueError("工作簿解压后过大")
        book = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except (BadZipFile, KeyError, OSError, ValueError, ParseError) as exc:
        raise ValueError("无效或过大的 xlsx 文件") from exc
    rows, sheet_headers = [], {}
    try:
        if len(book.sheetnames) > 40:
            raise ValueError("当前最多支持 40 张工作表")
        for sheet in book:
            if (sheet.max_row or 0) > 30001 or (sheet.max_column or 0) > 60:
                raise ValueError("每表最多 30000 条记录、60 列")
            sheet.reset_dimensions()
            iterator = sheet.iter_rows(values_only=True)
            headers = [cell(v) for v in next(iterator, ())]
            if len(headers) > 60:
                raise ValueError(f"{sheet.title} 超过 60 列")
            if any(h.startswith("=") for h in headers):
                raise ValueError(f"{sheet.title} 表头包含公式")
            if len(set(h for h in headers if h)) != len([h for h in headers if h]):
                raise ValueError(f"{sheet.title} 存在重复表头")
            sheet_headers[sheet.title] = [h for h in headers if h]
            for index, values in enumerate(iterator, 2):
                if len(values) > 60:
                    raise ValueError(f"{sheet.title} 第 {index} 行超过 60 列")
                if not any(v is not None and v != "" for v in values):
                    continue
                if any(
                    v is not None and (i >= len(headers) or not headers[i])
                    for i, v in enumerate(values)
                ):
                    raise ValueError(f"{sheet.title} 第 {index} 行存在没有表头的数据")
                raw = {
                    h: v.isoformat() if isinstance(v, (date, datetime)) else v
                    for h, v in zip(headers, values, strict=False)
                    if h
                }
                if any(isinstance(v, str) and v.startswith("=") for v in raw.values()):
                    raise ValueError(f"{sheet.title} 第 {index} 行含公式，请导出为值后导入")
                rows.append({"sheet": sheet.title, "row": index, "values": raw})
                if len(rows) > 30000:
                    raise ValueError("有效记录不能超过 30000 行")
    except ParseError as exc:
        raise ValueError("无效或过大的 xlsx 文件") from exc
    finally:
        book.close()
    return _WorkbookRows(rows, sheet_headers), sheet_headers


class _WorkbookRows(list):
    """A normal JSON-serializable list retaining empty-sheet metadata in memory."""

    def __init__(self, rows, headers):
        super().__init__(rows)
        self.headers = headers


def _input_rows(rows, headers=None):
    if isinstance(rows, tuple) and len(rows) == 2:
        rows, supplied = rows
        if headers is None:
            headers = supplied
    if headers is None:
        headers = getattr(rows, "headers", {})
    if not isinstance(rows, list) or len(rows) > 30000 or not isinstance(headers, dict):
        raise ValueError("记录或表头结构无效，有效记录最多 30000 行")
    columns = {}
    for name, names in headers.items():
        if (
            not isinstance(name, str)
            or not name.strip()
            or not isinstance(names, list)
            or len(names) > 60
            or any(not isinstance(n, str) or not n.strip() for n in names)
            or len(set(names)) != len(names)
        ):
            raise ValueError("表头结构无效")
        columns[name] = dict.fromkeys(names)
    normalized, seen = [], set()
    for row in rows:
        if (
            not isinstance(row, dict)
            or not isinstance(row.get("sheet"), str)
            or not row["sheet"].strip()
            or type(row.get("row")) is not int
            or row["row"] < 2
            or not isinstance(row.get("values"), dict)
        ):
            raise ValueError("来源记录结构无效")
        source = (row["sheet"], row["row"])
        if source in seen:
            raise ValueError("来源工作表和行号重复")
        seen.add(source)
        raw = {}
        for key, value in row["values"].items():
            if not isinstance(key, str) or not key.strip():
                raise ValueError("来源列名无效")
            if isinstance(value, (date, datetime)):
                value = value.isoformat()
            if (value is not None and type(value) not in (str, int, float, bool)) or (
                isinstance(value, float) and not math.isfinite(value)
            ):
                raise ValueError("来源值必须是有限数或 JSON 基本类型")
            raw[key] = value
        if row["sheet"] in headers and not set(raw) <= set(headers[row["sheet"]]):
            raise ValueError("来源记录包含表头之外的列")
        columns.setdefault(row["sheet"], {}).update(dict.fromkeys(raw))
        normalized.append({"sheet": row["sheet"], "row": row["row"], "values": raw})
    if len(columns) > 40 or any(len(c) > 60 for c in columns.values()):
        raise ValueError("最多支持 40 张工作表、每表 60 列")
    return normalized, {name: list(names) for name, names in columns.items()}


def group_rows(rows):
    if isinstance(rows, tuple) and len(rows) == 2:
        rows = rows[0]
    grouped = defaultdict(list)
    for r in rows:
        grouped[r["sheet"]].append(r)
    return grouped


def _join_candidates(indexes):
    candidates, comparisons, work = [], 0, 0
    # Hints are intentionally bounded, not an exhaustive schema search.
    for source_sheet, source_column, source in indexes:
        for target_sheet, target_column, target in indexes:
            if source_sheet == target_sheet or len(target) < len(source) * 0.8:
                continue
            comparisons += 1
            work += min(len(source), len(target))
            if comparisons > 20000 or work > 2000000:
                return candidates
            matched = len(source & target)
            if matched / len(source) >= 0.8:
                candidates.append(
                    {
                        "sourceSheet": source_sheet,
                        "sourceColumn": source_column,
                        "targetSheet": target_sheet,
                        "targetColumn": target_column,
                        "matched": matched,
                        "total": len(source),
                        "matchRate": matched / len(source),
                    }
                )
                if len(candidates) == 40:
                    return candidates
    return candidates


def profile_rows(rows, headers=None):
    rows, headers = _input_rows(rows, headers)
    grouped, sheets, indexes = group_rows(rows), [], []
    for name, names in headers.items():
        records, columns = grouped.get(name, []), []
        for column in names:
            values = [r["values"].get(column) for r in records if cell(r["values"].get(column))]
            counts = Counter(cell(v) for v in values)
            numeric = bool(values) and all(type(v) in (int, float) for v in values)
            kind = "number" if numeric else "text"
            if values and not numeric:
                try:
                    for v in values:
                        convert(v, "date")
                    kind = "date"
                except (ValueError, TypeError, OverflowError):
                    pass
            columns.append(
                {
                    "name": column,
                    "dataType": kind,
                    "nonNull": len(values),
                    "uniqueCount": len(counts),
                    "unique": bool(records) and len(counts) == len(records),
                    "samples": [v[:160] for v in list(counts)[:3]],
                }
            )
            if len(counts) >= 2:
                indexes.append((name, column, set(counts)))
        sheets.append({"name": name, "rows": len(records), "columns": columns})
    return {"rowCount": len(rows), "sheets": sheets, "joinCandidates": _join_candidates(indexes)}


def convert(value, kind):
    if kind not in ("text", "number", "date"):
        raise ValueError("未知属性类型")
    if value is not None and not isinstance(value, (str, int, float, bool, date, datetime)):
        raise ValueError("属性值必须是基本类型")
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("非有限数")
    if not cell(value):
        return None
    if kind == "number":
        if isinstance(value, (bool, date, datetime)):
            raise ValueError("无效数值类型")
        try:
            parsed = Decimal(cell(value))
            result = float(parsed)
            if not math.isfinite(result):
                raise ValueError("非有限数")
            if parsed and result == 0:
                raise ValueError("数值下溢")
            # Integers (including -0) have one canonical representation in entity keys.
            return int(parsed) if parsed == parsed.to_integral_value() else result
        except (InvalidOperation, OverflowError) as exc:
            raise ValueError("无效数值") from exc
    if kind == "date":
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        text = cell(value)
        if len(text) < 10 or text[4] != "-" or text[7] != "-":
            raise ValueError("日期必须使用 ISO 格式")
        if len(text) == 10:
            return date.fromisoformat(text).isoformat()
        return datetime.fromisoformat(text).date().isoformat()
    return cell(value)


MAX_NODES = 60000
MAX_EDGES = 100000


class _Report:
    def __init__(self):
        self.errors = 0
        self.warnings = 0
        self.issues = []
        self.issue_groups = {}
        self.groups_truncated = False
        self.metrics = []
        self.reviewed = True
        self.config_hash = ""
        self.class_counts = {}

    def issue(self, severity, item, message):
        if severity == "error":
            self.errors += 1
        else:
            self.warnings += 1
        if len(self.issues) < 200:
            self.issues.append({"severity": severity, "item": item, "message": message})
        key = (severity, item, message)
        if key in self.issue_groups:
            self.issue_groups[key]["count"] += 1
        elif len(self.issue_groups) < 500:
            self.issue_groups[key] = {
                "severity": severity,
                "item": item,
                "message": message,
                "count": 1,
            }
        else:
            self.groups_truncated = True

    def finish(self, nodes=(), edges=()):
        for node in nodes:
            self.class_counts[node["objectType"]] += 1
        return (
            {
                "errors": self.errors,
                "warnings": self.warnings,
                "issues": self.issues,
                "issueGroups": list(self.issue_groups.values()),
                "issueGroupsTruncated": self.groups_truncated,
                "relationships": self.metrics,
                "canPublish": not self.errors and self.reviewed,
                "configHash": self.config_hash,
                "counts": {
                    "nodes": len(nodes),
                    "edges": len(edges),
                    "classes": len(self.class_counts),
                },
                "classCounts": self.class_counts,
            },
            list(nodes),
            list(edges),
        )


def _check_model(model, grouped, headers, report, require_review, model_only=False):
    from .rules import check_definitions

    check_definitions(model, report, require_review)
    targets = {
        f"{kind}:{item.id}"
        for kind, items in (
            ("class", model.classes),
            ("mapping", model.mappings),
            ("relationship", model.relationships),
        )
        for item in items
    }
    targets |= {f"property:{cls.id}:{p.id}" for cls in model.classes for p in cls.properties}
    for items, key, label in (
        (model.semantics.questions, "id", "业务问题"),
        (model.semantics.definitions, "target", "语义定义"),
    ):
        seen = set()
        for item in items:
            ident = getattr(item, key)
            if ident in seen:
                report.issue("error", ident, f"{label}标识重复")
            seen.add(ident)
            if key == "target" and ident not in targets:
                report.issue("error", ident, "语义定义引用不存在的建模项")
            if key == "target" and not item.confirmed:
                report.issue("warning", ident, "语义口径尚未业务确认；发布后仍标记为待确认")
    classes = {c.id: c for c in model.classes}
    properties = {c.id: {p.id for p in c.properties} for c in model.classes}
    for collection in (
        model.classes,
        [] if model_only else model.mappings,
        model.relationships,
        [] if model_only else model.instances,
        [] if model_only else model.links,
    ):
        seen = set()
        for item in collection:
            if item.id in seen:
                report.issue("error", item.id, "标识重复")
            seen.add(item.id)
            if hasattr(item, "accepted") and not item.accepted:
                report.reviewed = False
                report.issue("error" if require_review else "warning", item.id, "尚未审核")
    for cls in model.classes:
        if len(properties[cls.id]) != len(cls.properties):
            report.issue("error", cls.id, "属性标识重复")
        if len({p.name for p in cls.properties}) != len(cls.properties):
            report.issue("error", cls.id, "属性显示名称重复")
        if len(set(cls.identity)) != len(cls.identity):
            report.issue("error", cls.id, "身份属性重复")
        if not {cls.labelProperty, *cls.identity} <= properties[cls.id]:
            report.issue("error", cls.id, "身份或标签引用不存在的属性")
        if not cls.identity:
            report.issue(
                "warning", cls.id, "没有 identity，使用来源行标识，逐行不合并且不能跨文件对齐"
            )
    for cls in model.classes:
        if cls.parentClassId is None:
            continue
        parent = classes.get(cls.parentClassId)
        if parent is None:
            report.issue("error", cls.id, "继承引用不存在的父类")
            continue
        if parent.id == cls.id:
            report.issue("error", cls.id, "父类不能是自身")
            continue
        if not cls.inheritanceReason:
            report.issue("error", cls.id, "继承缺少 is-a 业务依据，请补充并人工审核父类关系")
        parent_props = {p.id: p.dataType for p in parent.properties}
        child_props = {p.id: p.dataType for p in cls.properties}
        if any(child_props.get(pid) != kind for pid, kind in parent_props.items()):
            report.issue("error", cls.id, "子类必须显式保留父类全部属性的标识及数据类型")
        if not cls.identity:
            report.issue(
                "error", cls.id, "父类与子类继承必须明确非空业务身份，不能按来源行猜测同一实体"
            )
        if cls.identity != parent.identity:
            report.issue("error", cls.id, "子类必须与父类保持相同且顺序一致的业务身份属性")
        ancestors = {cls.id}
        cursor = cls.parentClassId
        while cursor is not None and cursor in classes:
            if cursor in ancestors:
                report.issue("error", cls.id, "父类继承存在循环")
                break
            ancestors.add(cursor)
            if len(ancestors) > 16:
                report.issue("error", cls.id, "父类继承链超过平台支持的 16 层")
                break
            cursor = classes[cursor].parentClassId
    for mapping in [] if model_only else model.mappings:
        cls = classes.get(mapping.classId)
        if cls is None:
            report.issue("error", mapping.id, "映射引用不存在的类")
            continue
        if not set(mapping.bindings) <= properties[cls.id]:
            report.issue("error", mapping.id, "绑定引用不存在的属性")
        if not {cls.labelProperty, *cls.identity} <= set(mapping.bindings):
            report.issue("error", mapping.id, "绑定必须包含 identity 和 labelProperty")
        if mapping.sheet not in headers:
            report.issue("error", mapping.id, "来源工作表不存在")
        elif not set(mapping.bindings.values()) <= set(headers[mapping.sheet]):
            report.issue("error", mapping.id, "绑定列不存在")
        elif not grouped.get(mapping.sheet):
            report.issue("warning", mapping.id, "来源为空表，保留类定义但不生成实例")
    for relation in model.relationships:
        if relation.source not in classes or relation.target not in classes:
            report.issue("error", relation.id, "关系引用不存在的类")
            continue
        joins = [(j.sourceProperty, j.targetProperty) for j in relation.joins]
        if len(set(joins)) != len(joins):
            report.issue("error", relation.id, "关联条件重复")
        if any(
            s not in properties[relation.source] or t not in properties[relation.target]
            for s, t in joins
        ):
            report.issue("error", relation.id, "joins 引用不存在的属性")
    return classes


def _values(raw, cls, report, item):
    result = {}
    for prop in cls.properties:
        try:
            result[prop.id] = convert(raw.get(prop.id), prop.dataType)
        except (ValueError, TypeError, OverflowError):
            report.issue("error", item, f"属性 {prop.id} 不符合 {prop.dataType} 类型")
            result[prop.id] = None
    return result


def _add_manual_instances(model, entities, classes, report):
    for manual in model.instances:
        cls = classes.get(manual.classId)
        if cls is None:
            report.issue("error", manual.id, "人工实例引用不存在的类")
            continue
        if not set(manual.values) <= {p.id for p in cls.properties}:
            report.issue("error", manual.id, "人工实例引用不存在的属性")
        properties = _values(manual.values, cls, report, manual.id)
        identity = [properties[p] for p in cls.identity]
        if any(v is None or v == "" for v in identity):
            report.issue("error", manual.id, "人工实例 identity 身份属性存在空值")
        key = f"entity:{cls.id}:manual:{manual.id}"
        if len(entities) >= MAX_NODES:
            report.issue("error", "instances", "实例超过 60000 个")
            return
        entities[key] = {
            "key": key,
            "classId": cls.id,
            "identity": identity,
            "properties": properties,
            "sources": [],
            "sourceValues": dict(manual.values),
            "conflicts": set(),
            "manual": True,
            "edited": False,
        }


def _apply_instance_edits(model, entities, classes, report, isolate_conflicts=False):
    seen = set()
    for edit in model.instanceEdits:
        if edit.key in seen:
            report.issue("error", "instanceEdits", "实例编辑 key 重复")
        seen.add(edit.key)
        entity = entities.get(edit.key)
        if entity is None:
            report.issue("error", "instanceEdits", "编辑引用已经不存在的实例，请重新审核映射")
            continue
        cls = classes[entity["classId"]]
        if not set(edit.values) <= set(entity["properties"]):
            report.issue("error", edit.key, "实例编辑引用不存在的属性")
        if set(edit.values) & set(cls.identity):
            report.issue(
                "error", edit.key, "不允许编辑 identity 身份字段，请修改来源并重新审核映射"
            )
            continue
        converted = _values(edit.values, cls, report, edit.key)
        entity["properties"].update({k: converted[k] for k in edit.values if k in converted})
        entity["conflicts"].difference_update(edit.values)
        entity["edited"] = bool(edit.values)
        entity["excluded"] = edit.excluded
    for entity in entities.values():
        if entity["conflicts"] and not isolate_conflicts:
            cls = classes[entity["classId"]]
            locations = "、".join(f"{row['sheet']}:{row['row']}" for row in entity["sources"][:5])
            report.issue(
                "error",
                entity["key"],
                "同身份实例非空属性冲突："
                + "、".join(sorted(entity["conflicts"]))
                + f"；对象 {cls.name}，粒度：{cls.grain or '未说明'}；来源行：{locations or '手工维护'}。"
                + "请核对数据错误、版本或粒度差异后处理，不应直接覆盖冲突值。",
            )


def _json(value):
    return json.dumps(value, ensure_ascii=False, allow_nan=False)


def _node(key, cls, identity, properties, sources, manual=False, edited=False, source_values=None):
    first = (
        sources[0] if sources else {"sheet": "手工维护", "row": 0, "values": source_values or {}}
    )
    payload = {
        "sheet": first["sheet"],
        "row": first["row"],
        "sourceValues": first["values"],
        "values": {p.name: properties[p.id] for p in cls.properties},
        "properties": properties,
        "sources": sources,
        "manual": manual,
        "edited": edited,
    }
    return {
        "key": key,
        "kind": cls.name,
        "label": cell(properties[cls.labelProperty]),
        "objectType": cls.id,
        "identity": _json(identity),
        "sheet": first["sheet"],
        "row": first["row"],
        "payload": _json(payload),
    }


def _join_key(properties, keys):
    values = tuple(properties[k] for k in keys)
    return values if all(v is not None and v != "" for v in values) else None


def _check_link_edits(model, entities, report):
    relations = {r.id: r for r in model.relationships}
    edits, seen = defaultdict(list), set()
    for link in model.links:
        triple = (link.relationId, link.source, link.target)
        if triple in seen:
            report.issue("error", link.id, "相同关系和端点的边编辑重复")
        seen.add(triple)
        relation = relations.get(link.relationId)
        if relation is None:
            report.issue("error", link.id, "边引用不存在的关系")
            continue
        source, target = entities.get(link.source), entities.get(link.target)
        if source is None or target is None:
            report.issue("error", link.id, "边引用不存在的实例端点")
        elif source["classId"] != relation.source or target["classId"] != relation.target:
            report.issue("error", link.id, "边端点与关系的来源或目标类不匹配")
        elif source.get("excluded") or target.get("excluded"):
            report.issue("warning", link.id, "边端点实例已排除，该边自动移除")
        else:
            edits[link.relationId].append(link)
    return edits


def _compile_relations(model, entities, link_edits, report):
    by_class = defaultdict(list)
    for entity in entities.values():
        by_class[entity["classId"]].append(entity)
    edges = []
    for relation in model.relationships:
        source_keys = [j.sourceProperty for j in relation.joins]
        target_keys = [j.targetProperty for j in relation.joins]
        index = defaultdict(set)
        for target in by_class[relation.target]:
            key = _join_key(target["properties"], target_keys)
            if key is not None:
                index[key].add(target["key"])
        additions, removals = defaultdict(set), defaultdict(set)
        for link in link_edits.get(relation.id, []):
            (removals if link.excluded else additions)[link.source].add(link.target)
        pairs, matched, count, multiple = [], 0, 0, False
        for source in by_class[relation.source]:
            targets = index.get(_join_key(source["properties"], source_keys), frozenset())
            removed = removals.get(source["key"], frozenset())
            extra = additions.get(source["key"], frozenset()) - targets
            size = len(targets) - len(targets & removed) + len(extra)
            count += size
            matched += size > 0
            multiple |= size > 1
            # Count with indexed set sizes BEFORE expanding any Cartesian product.
            if len(edges) + count > MAX_EDGES:
                report.issue("error", relation.id, "关系超过 100000 条，请收窄关联条件")
                return []
            pairs.append((source["key"], targets, removed, extra))
        total = len(by_class[relation.source])
        report.metrics.append(
            {
                "id": relation.id,
                "matched": matched,
                "total": total,
                "edges": count,
                "matchRate": matched / total if total else 0,
            }
        )
        if matched < total:
            report.issue(
                "warning", relation.id, f"{total - matched} 个实例无法关联，不生成对应连线"
            )
        if multiple:
            report.issue("warning", relation.id, "目标属性不唯一，将保留一对多关系")
        for source, targets, removed, extra in pairs:
            edges.extend(
                {
                    "source": source,
                    "target": target,
                    "label": relation.name,
                    "relationId": relation.id,
                }
                for target in sorted((targets - removed) | extra)
            )
    return edges


def validate_ontology(value, require_review=True):
    """Validate publishable definitions without interpreting or writing source instances."""
    report = _Report()
    try:
        model = Draft.model_validate(value.model_dump() if isinstance(value, Draft) else value)
    except (ValidationError, ValueError, TypeError):
        report.issue("error", "draft", "候选本体结构无效，请检查字段、类型和容量限制")
    else:
        report.config_hash = digest(model.model_dump())
        from .scope import publication_scope

        scoped = publication_scope(model.model_dump())
        if not scoped["classes"]:
            report.issue("error", "draft", "请至少选择一个对象类型用于本次发布")
        else:
            model = Draft.model_validate(scoped)
            report.class_counts = {c.id: 0 for c in model.classes}
            _check_model(model, {}, {}, report, require_review, model_only=True)
    result = report.finish()[0]
    result["scope"] = "ontology"
    if not result["errors"]:
        result["counts"].update(
            properties=sum(len(c.properties) for c in model.classes),
            relationships=len(model.relationships),
        )
    return result


def _evaluate(value, rows, require_review, quarantine=None):
    report = _Report()
    try:
        # Revalidate model instances as well as dictionaries; do not trust model_copy updates.
        model = Draft.model_validate(value.model_dump() if isinstance(value, Draft) else value)
    except (ValidationError, ValueError, TypeError):
        report.issue(
            "error", "draft", "草稿结构无效：字段缺失、非法字段、类型或数量不符合 schemaVersion 2"
        )
        return report.finish()
    report.config_hash = digest(model.model_dump())
    from .scope import publication_scope

    scoped = publication_scope(model.model_dump())
    if not scoped["classes"]:
        report.issue("error", "draft", "请至少选择一个对象类型用于本次发布")
        return report.finish()
    model = Draft.model_validate(scoped)
    report.class_counts = {c.id: 0 for c in model.classes}
    try:
        rows, headers = _input_rows(rows)
    except (ValueError, TypeError, OverflowError):
        report.issue("error", "rows", "来源记录或表头无效：请检查类型、重复行、非有限数和容量限制")
        return report.finish()
    grouped = group_rows(rows)
    classes = _check_model(model, grouped, headers, report, require_review)
    if report.errors:
        return report.finish()
    entities = {}
    invalid_keys = set()

    def defer(kind, item, reason, sources=(), **extra):
        quarantine.append(
            {
                "kind": kind,
                "item": item,
                "reason": reason,
                "sources": [{"sheet": r["sheet"], "row": r["row"]} for r in sources],
                **extra,
            }
        )

    for mapping in model.mappings:
        cls = classes[mapping.classId]
        for row in grouped.get(mapping.sheet, []):
            row_report = report if quarantine is None else _Report()
            properties = _values(
                {p: row["values"].get(c) for p, c in mapping.bindings.items()},
                cls,
                row_report,
                mapping.id,
            )
            identity = [properties[p] for p in cls.identity]
            if cls.identity and any(v is None or v == "" for v in identity):
                if quarantine is None:
                    report.issue("error", mapping.id, "identity 身份属性存在空值")
                else:
                    defer("row", mapping.id, "identity 身份属性存在空值", [row], classId=cls.id)
                continue
            key = (
                f"entity:{cls.id}:" + digest(identity)[:24]
                if cls.identity
                else f"entity:{cls.id}:row:{mapping.id}:{row['row']}"
            )
            source_id = (row["sheet"], row["row"])
            if quarantine is not None and row_report.errors:
                invalid_keys.add(key)
                defer(
                    "row",
                    mapping.id,
                    "；".join(i["message"] for i in row_report.issues),
                    [row],
                    classId=cls.id,
                )
                continue
            if key not in entities:
                entities[key] = {
                    "key": key,
                    "classId": cls.id,
                    "identity": identity,
                    "properties": properties,
                    "sources": [row],
                    "sourceIds": {source_id},
                    "conflicts": set(),
                    "edited": False,
                    "manual": False,
                }
            else:
                entity = entities[key]
                if source_id not in entity["sourceIds"]:
                    entity["sources"].append(row)
                    entity["sourceIds"].add(source_id)
                for prop, incoming in properties.items():
                    current = entity["properties"][prop]
                    if current is None:
                        entity["properties"][prop] = incoming
                    elif incoming is not None and current != incoming:
                        entity["conflicts"].add(prop)
            if len(entities) > MAX_NODES:
                report.issue("error", "instances", "实例超过 60000 个")
                return report.finish()
    _add_manual_instances(model, entities, classes, report)
    if quarantine is None:
        _apply_instance_edits(model, entities, classes, report)
    else:
        # Invalid edit references or identity edits remain configuration blockers.
        _apply_instance_edits(model, entities, classes, report, isolate_conflicts=True)
        for key, entity in entities.items():
            if entity["conflicts"] or key in invalid_keys:
                reason = (
                    "同身份实例非空属性冲突：" + "、".join(sorted(entity["conflicts"]))
                    if entity["conflicts"]
                    else "同身份的其他来源记录存在类型错误，整个身份组暂缓"
                )
                defer(
                    "entity",
                    key,
                    reason,
                    entity["sources"],
                    classId=entity["classId"],
                    conflicts=sorted(entity["conflicts"]),
                    identity=entity["identity"],
                )
                entity["excluded"] = True
    link_edits = _check_link_edits(model, entities, report)
    if report.errors:
        return report.finish()
    active = {k: e for k, e in entities.items() if not e.get("excluded")}
    from .rules import check_unique_values

    if quarantine is None:
        check_unique_values(model, active, report)
    else:
        unique_report = _Report()
        invalid_unique = check_unique_values(model, active, unique_report)
        for key, entity in list(active.items()):
            if key in invalid_unique:
                defer(
                    "entity",
                    key,
                    "已确认替代唯一键为空或重复，相关身份组暂缓",
                    entity["sources"],
                    classId=entity["classId"],
                )
                entity["excluded"] = True
                del active[key]
    if len(active) != len(entities):
        report.issue("warning", "instanceEdits", "已排除实例，其关联边自动移除")
    identity_index = {
        (entity["classId"], digest(entity["identity"])): entity
        for entity in active.values()
        if classes[entity["classId"]].identity
    }
    duplicate_keys = set()
    for entity in active.values():
        cls = classes[entity["classId"]]
        ancestor = cls.parentClassId
        while ancestor is not None:
            if (ancestor, digest(entity["identity"])) in identity_index:
                reason = f"同一业务身份同时生成父类 {ancestor} 和子类 {cls.id} 实例；请只映射到最具体的类"
                if quarantine is None:
                    report.issue("error", entity["key"], reason)
                else:
                    duplicate_keys.update(
                        [
                            entity["key"],
                            identity_index[(ancestor, digest(entity["identity"]))]["key"],
                        ]
                    )
            ancestor = classes[ancestor].parentClassId
    if report.errors:
        return report.finish()
    for key in sorted(duplicate_keys):
        entity = active.pop(key)
        defer(
            "entity",
            key,
            "父子类重复生成同一身份，相关身份组暂缓",
            entity["sources"],
            classId=entity["classId"],
        )
        entity["excluded"] = True
    if quarantine is not None:
        # Remove manual links to every deferred endpoint, even after unique/inheritance checks.
        for link in model.links:
            if not link.excluded and (link.source not in active or link.target not in active):
                defer("relationship", link.relationId, "人工连线端点暂缓，不写入该连线", count=1)
        link_edits = {
            rid: [link for link in links if link.source in active and link.target in active]
            for rid, links in link_edits.items()
        }
    edges = _compile_relations(model, active, link_edits, report)
    from .rules import check_edge_cardinality

    if quarantine is None:
        check_edge_cardinality(model, edges, report)
    else:
        edge_report = _Report()
        check_edge_cardinality(model, edges, edge_report)
        bad_relations = {i["item"] for i in edge_report.issues if i["severity"] == "error"}
        for rid in sorted(bad_relations):
            count = sum(e["relationId"] == rid for e in edges)
            defer("relationship", rid, "实际连线违反已确认基数，本次暂缓该关系", count=count)
        edges = [e for e in edges if e["relationId"] not in bad_relations]
        for metric in report.metrics:
            if metric["total"] > metric["matched"]:
                defer(
                    "relationship",
                    metric["id"],
                    "缺少可用目标或关联字段为空，不写入对应连线",
                    count=metric["total"] - metric["matched"],
                )
    if report.errors:
        return report.finish()
    nodes = [
        _node(
            e["key"],
            classes[e["classId"]],
            e["identity"],
            e["properties"],
            e["sources"],
            e["manual"],
            e["edited"],
            e.get("sourceValues"),
        )
        for e in active.values()
    ]
    return report.finish(nodes, edges)


def validate_draft(value, rows, require_review=True):
    """Validate and preview through the same bounded compiler used for publication."""
    return _evaluate(value, rows, require_review)[0]


def compile_graph(value, rows, require_review=True):
    report, nodes, edges = _evaluate(value, rows, require_review)
    if report["errors"]:
        raise ValueError("配置未通过审核或校验，请先查看 validate_draft 的 issues")
    return nodes, edges
